"""DOCX handler.

Works at the paragraph level: collects each paragraph's full text, anonymizes,
and writes the result back into the first run (clearing the rest). This catches
entities that span run boundaries (a common Word quirk where a single name gets
split across runs because of cursor positions) but sacrifices sub-paragraph
formatting variation in modified paragraphs. Unmodified paragraphs are untouched.
"""
from __future__ import annotations

from pathlib import Path

from .base import FormatHandler


class DocxHandler(FormatHandler):
    extensions = (".docx",)

    def process(self, input_path: Path, output_path: Path, engine) -> dict:
        try:
            from docx import Document
        except ImportError as e:
            raise RuntimeError(
                "python-docx is required for DOCX processing. "
                "Install with: pip install python-docx"
            ) from e

        doc = Document(str(input_path))
        findings: list[dict] = []

        def process_paragraph(p) -> None:
            text = p.text
            if not text or not text.strip():
                return
            anonymized, results = engine.anonymize(text)
            findings.extend(
                {"entity_type": r.entity_type, "score": round(r.score, 3)} for r in results
            )
            if anonymized != text and p.runs:
                p.runs[0].text = anonymized
                for run in p.runs[1:]:
                    run.text = ""

        for paragraph in doc.paragraphs:
            process_paragraph(paragraph)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        process_paragraph(paragraph)

        # Also walk headers/footers
        for section in doc.sections:
            for hf in (section.header, section.footer):
                for paragraph in hf.paragraphs:
                    process_paragraph(paragraph)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return {"format": "docx", "findings": findings}
